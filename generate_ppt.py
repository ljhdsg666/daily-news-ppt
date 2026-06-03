import os
import io
import json
import requests
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from openai import OpenAI
from datetime import datetime
import feedparser
from PIL import Image, ImageDraw

# ---------- 配置 ----------
DEEPSEEK_API_KEY = os.environ['DEEPSEEK_API_KEY']
WECOM_WEBHOOK = os.environ['WECOM_WEBHOOK']

client = OpenAI(
    api_key=DEEPSEEK_API_KEY,
    base_url="https://api.deepseek.com"
)

# 专业配色方案（深蓝+金）
PRIMARY_COLOR = RGBColor(0x0D, 0x2C, 0x54)   # 深蓝
ACCENT_COLOR = RGBColor(0xD4, 0xA8, 0x3C)    # 金色
BG_LIGHT = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(255, 255, 255)

# ---------- 1. 抓取新闻 ----------
def fetch_news():
    query = "wall street site:usnews.com"
    url = f"https://news.google.com/rss/search?q={requests.utils.quote(query)}&hl=en-US&gl=US&ceid=US:en"
    feed = feedparser.parse(url)
    articles = []
    for entry in feed.entries[:8]:   # 取8条做多页
        articles.append({
            "title": entry.title,
            "link": entry.link
        })
    return articles

# ---------- 2. AI 生成结构化PPT方案 ----------
def generate_ppt_blueprint(articles, date_str):
    """让DeepSeek直接输出一个完整的PPT设计方案（JSON）"""
    news_list = "\n".join([f"{i+1}. {a['title']}" for i, a in enumerate(articles)])
    prompt = f"""你是一个顶级咨询公司的演示文稿设计师。请根据以下今日华尔街新闻，设计一份专业PPT的完整方案。
要求：
- 输出严格JSON格式，不要有任何额外文字。
- 包含封面页、目录页、3~5条重点新闻的详细页、一个总结页。
- 每条新闻页需要包含：标题、3~4个要点（每个要点15字左右）、图片搜索关键词（英文）、可选图表（如果涉及数据）。
- 图表用 chart 字段，格式为 {{"type": "bar"或"line", "categories": ["类1","类2"], "values": [数值1,数值2]}}，没有数据则不要。
- 总结页用3个精炼的句子概括今日重点。

JSON结构：
{{
  "slides": [
    {{"type": "cover", "title": "每日华尔街简报", "subtitle": "{date_str}"}},
    {{"type": "toc", "items": ["新闻标题1", "新闻标题2", ...]}},
    {{"type": "content", "title": "...", "points": ["要点1", "要点2", "要点3"], "image_keyword": "finance chart", "chart": {{...}} }},
    ...,
    {{"type": "summary", "text": "总结内容..."}}
  ]
}}

今日新闻：
{news_list}

请直接输出JSON。"""
    for _ in range(3):  # 最多尝试3次解析JSON
        try:
            response = client.chat.completions.create(
                model="deepseek-chat",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.5,
            )
            raw = response.choices[0].message.content.strip()
            # 去掉可能的markdown代码块标记
            if raw.startswith("```json"):
                raw = raw[7:]
            if raw.endswith("```"):
                raw = raw[:-3]
            blueprint = json.loads(raw)
            return blueprint.get("slides", [])
        except Exception as e:
            print(f"JSON解析失败，重试... {e}")
            continue
    # 兜底：返回简单结构
    print("AI生成PPT方案失败，使用默认结构")
    return [{"type": "cover", "title": "每日华尔街简报", "subtitle": date_str}]

# ---------- 3. 图片获取 ----------
def get_image(keyword):
    try:
        url = f"https://picsum.photos/800/600?random&{requests.utils.quote(keyword)}"
        resp = requests.get(url, timeout=8)
        if resp.status_code == 200 and 'image' in resp.headers.get('Content-Type',''):
            img = Image.open(io.BytesIO(resp.content))
            img_io = io.BytesIO()
            img.save(img_io, format='PNG')
            img_io.seek(0)
            return img_io
    except Exception as e:
        print(f"图片下载失败: {e}")
    # 占位渐变图
    img = Image.new('RGB', (800,600), (13,44,84))
    draw = ImageDraw.Draw(img)
    for i in range(600):
        color = (13 + i//15, 44 + i//20, 84 + i//25)
        draw.line([(0,i), (800,i)], fill=color)
    img_io = io.BytesIO()
    img.save(img_io, format='PNG')
    img_io.seek(0)
    return img_io

# ---------- 4. 生成专业PPT ----------
def create_ppt_from_blueprint(slides_data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 主题色设置（作为背景或形状）
    def add_bg(slide, color=WHITE):
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    for slide_data in slides_data:
        stype = slide_data.get("type")

        # --- 封面页 ---
        if stype == "cover":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            # 背景图
            img = get_image("wall street skyscraper")
            slide.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
            # 深色遮罩
            shade = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
            shade.fill.solid()
            shade.fill.fore_color.rgb = RGBColor(0,0,0)
            shade.fill.fore_color.brightness = -0.5
            shade.line.fill.background()
            # 标题
            txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10), Inches(1.5))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = slide_data.get("title", "每日简报")
            p.font.size = Pt(54)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
            # 副标题
            txBox2 = slide.shapes.add_textbox(Inches(1.5), Inches(4.0), Inches(10), Inches(0.8))
            tf2 = txBox2.text_frame
            p2 = tf2.paragraphs[0]
            p2.text = slide_data.get("subtitle", "")
            p2.font.size = Pt(24)
            p2.font.color.rgb = ACCENT_COLOR
            p2.alignment = PP_ALIGN.CENTER

        # --- 目录页 ---
        elif stype == "toc":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_bg(slide, WHITE)
            # 左侧色条
            bar = slide.shapes.add_shape(1, 0, 0, Inches(0.8), prs.slide_height)
            bar.fill.solid()
            bar.fill.fore_color.rgb = PRIMARY_COLOR
            bar.line.fill.background()
            # 标题
            txBox = slide.shapes.add_textbox(Inches(1.5), Inches(0.8), Inches(10), Inches(0.8))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = "今日概览"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = PRIMARY_COLOR
            # 条目
            items = slide_data.get("items", [])
            y_start = 2.2
            for i, item in enumerate(items[:5]):
                txBox = slide.shapes.add_textbox(Inches(2), Inches(y_start + i*0.9), Inches(9), Inches(0.7))
                tf = txBox.text_frame
                p = tf.paragraphs[0]
                p.text = f"{i+1}.  {item}"
                p.font.size = Pt(22)
                p.font.color.rgb = RGBColor(50,50,50)

        # --- 内容页（图文+要点，可带图表）---
        elif stype == "content":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_bg(slide, BG_LIGHT)
            # 顶部标题栏
            header = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.2))
            header.fill.solid()
            header.fill.fore_color.rgb = PRIMARY_COLOR
            header.line.fill.background()
            txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11), Inches(0.8))
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = slide_data.get("title", "")
            p.font.size = Pt(32)
            p.font.bold = True
            p.font.color.rgb = WHITE

            # 图片
            img_keyword = slide_data.get("image_keyword", "finance")
            img = get_image(img_keyword)
            slide.shapes.add_picture(img, Inches(0.5), Inches(1.6), Inches(4.2), Inches(4.2))

            # 要点文字
            points = slide_data.get("points", [])
            txBox = slide.shapes.add_textbox(Inches(5.2), Inches(1.6), Inches(7), Inches(4))
            tf = txBox.text_frame
            tf.word_wrap = True
            for i, point in enumerate(points):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = f"▸ {point}"
                p.font.size = Pt(20)
                p.font.color.rgb = RGBColor(30,30,30)
                p.space_after = Pt(10)

            # 可选图表
            chart_data = slide_data.get("chart")
            if chart_data and chart_data.get("categories") and chart_data.get("values"):
                try:
                    chart_left = Inches(0.5)
                    chart_top = Inches(5.9)
                    chart_width = Inches(4.2)
                    chart_height = Inches(1.3)
                    chart_frame = slide.shapes.add_chart(
                        XL_CHART_TYPE.BAR_CLUSTERED if chart_data.get("type")=="bar" else XL_CHART_TYPE.LINE,
                        chart_left, chart_top, chart_width, chart_height
                    )
                    chart = chart_frame.chart
                    cat_data = CategoryChartData()
                    cat_data.categories = chart_data["categories"]
                    cat_data.add_series('', chart_data["values"])
                    chart.replace_data(cat_data)
                    chart.has_legend = False
                    chart.style = 2
                except Exception as e:
                    print(f"图表生成失败: {e}")

        # --- 总结页 ---
        elif stype == "summary":
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_bg(slide, PRIMARY_COLOR)
            txBox = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10), Inches(4))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = "总结与展望"
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = ACCENT_COLOR
            p.alignment = PP_ALIGN.CENTER
            summary_text = slide_data.get("text", "")
            p2 = tf.add_paragraph()
            p2.text = summary_text
            p2.font.size = Pt(24)
            p2.font.color.rgb = WHITE
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = Pt(30)

    ppt_filename = "daily_news.pptx"
    prs.save(ppt_filename)
    return ppt_filename

# ---------- 5. 发送 ----------
def send_to_wecom(file_path, webhook):
    key = webhook.split('key=')[-1]
    upload_url = f"https://qyapi.weixin.qq.com/cgi-bin/webhook/upload_media?key={key}&type=file"
    with open(file_path, 'rb') as f:
        files = {'media': f}
        r = requests.post(upload_url, files=files)
    media_id = r.json()['media_id']
    data = {"msgtype": "file", "file": {"media_id": media_id}}
    requests.post(webhook, json=data)

# ---------- 主流程 ----------
if __name__ == '__main__':
    articles = fetch_news()
    if not articles:
        print("没有抓到新闻，退出")
        exit(1)
    today = datetime.now().strftime("%Y年%m月%d日")
    slides = generate_ppt_blueprint(articles, today)
    ppt_file = create_ppt_from_blueprint(slides)
    send_to_wecom(ppt_file, WECOM_WEBHOOK)
    print("推送完成")
